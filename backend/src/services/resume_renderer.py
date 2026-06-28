from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor
from jinja2 import Template

from src.config import settings
from src.models.profile import CandidateProfile
from src.models.resume_template import DEFAULT_TEMPLATE, ResumeTemplateSettings


RESUME_JINJA = """{{ contact.name or 'Candidate' }}
{% for field in header_line %}{{ field }}{% if not loop.last %} | {% endif %}{% endfor %}

{% if summary and 'summary' in sections %}SUMMARY
{{ summary }}

{% endif %}"""


class ResumeRenderer:
    def __init__(self, template_settings: ResumeTemplateSettings | None = None):
        self.template = template_settings or DEFAULT_TEMPLATE

    def _build_data(self, profile: CandidateProfile, tailored: dict | None = None) -> dict:
        contact = profile.contact.model_dump()
        header_parts = []
        for field in self.template.header_fields:
            val = contact.get(field)
            if val:
                header_parts.append(str(val))

        experience = (tailored or {}).get("experience") or [e.model_dump() for e in profile.experience]
        return {
            "contact": contact,
            "header_line": header_parts,
            "summary": (tailored or {}).get("summary") or profile.summary or "",
            "experience": experience,
            "education": [e.model_dump() for e in profile.education],
            "skills": (tailored or {}).get("skills") or profile.skills or [],
            "certifications": profile.certifications or [],
            "sections": self.template.section_order,
            "bullet_style": self.template.bullet_style.value,
        }

    def render_text(self, profile: CandidateProfile, tailored: dict | None = None) -> str:
        data = self._build_data(profile, tailored)
        tmpl = self.template
        lines: list[str] = []
        lines.append(data["contact"].get("name") or "Candidate")
        if data["header_line"]:
            lines.append(" | ".join(data["header_line"]))
        lines.append("")
        for section_key in tmpl.section_order:
            if section_key == "summary" and data["summary"]:
                lines.extend(["SUMMARY", data["summary"], ""])
            elif section_key == "experience" and data["experience"]:
                lines.append("EXPERIENCE")
                for exp in data["experience"]:
                    title_line = f"{exp.get('title', '')} | {exp.get('company', '')}"
                    if exp.get("start_date"):
                        title_line += f" | {exp['start_date']}"
                        if exp.get("end_date"):
                            title_line += f" – {exp['end_date']}"
                    lines.append(title_line)
                    for bullet in exp.get("bullets") or []:
                        lines.append(f"{tmpl.bullet_style.value} {bullet}")
                lines.append("")
            elif section_key == "education" and data["education"]:
                lines.append("EDUCATION")
                for edu in data["education"]:
                    line = f"{edu.get('degree', '')} | {edu.get('institution', '')}"
                    if edu.get("graduation_date"):
                        line += f" | {edu['graduation_date']}"
                    lines.append(line)
                lines.append("")
            elif section_key == "skills" and data["skills"]:
                lines.extend(["SKILLS", ", ".join(data["skills"]), ""])
            elif section_key == "certifications" and data["certifications"]:
                lines.extend(["CERTIFICATIONS", ", ".join(data["certifications"]), ""])
        return "\n".join(lines).strip()

    def render_docx(
        self,
        profile: CandidateProfile,
        tailored: dict | None,
        output_path: Path,
        template_settings: ResumeTemplateSettings | None = None,
    ) -> Path:
        tmpl = template_settings or self.template
        data = self._build_data(profile, tailored)

        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(tmpl.margin_inches)
            section.bottom_margin = Inches(tmpl.margin_inches)
            section.left_margin = Inches(tmpl.margin_inches)
            section.right_margin = Inches(tmpl.margin_inches)

        style = doc.styles["Normal"]
        style.font.name = tmpl.font_name.value
        style.font.size = Pt(tmpl.font_size_body)

        name_para = doc.add_paragraph()
        name_run = name_para.add_run(data["contact"].get("name") or "Candidate")
        name_run.bold = True
        name_run.font.size = Pt(tmpl.font_size_name)
        name_run.font.name = tmpl.font_name.value

        if data["header_line"]:
            header_para = doc.add_paragraph(" | ".join(data["header_line"]))
            header_para.runs[0].font.size = Pt(tmpl.font_size_body)
            header_para.runs[0].font.name = tmpl.font_name.value
            header_para.runs[0].font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        for section_key in tmpl.section_order:
            if section_key == "experience":
                if not data["experience"]:
                    continue
                if tmpl.include_section_headers:
                    self._add_section_header(doc, "EXPERIENCE", tmpl)
                for exp in data["experience"]:
                    title_line = f"{exp.get('title', '')} | {exp.get('company', '')}"
                    if exp.get("start_date"):
                        title_line += f" | {exp['start_date']}"
                        if exp.get("end_date"):
                            title_line += f" – {exp['end_date']}"
                    p = doc.add_paragraph()
                    run = p.add_run(title_line)
                    run.bold = True
                    run.font.name = tmpl.font_name.value
                    run.font.size = Pt(tmpl.font_size_body)
                    for bullet in exp.get("bullets") or []:
                        bp = doc.add_paragraph()
                        bp.paragraph_format.left_indent = Inches(0.25)
                        bp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                        bp.paragraph_format.line_spacing = tmpl.line_spacing
                        br = bp.add_run(f"{tmpl.bullet_style.value} {bullet}")
                        br.font.name = tmpl.font_name.value
                        br.font.size = Pt(tmpl.font_size_body)
                continue

            if section_key == "education":
                if not data["education"]:
                    continue
                if tmpl.include_section_headers:
                    self._add_section_header(doc, "EDUCATION", tmpl)
                for edu in data["education"]:
                    line = f"{edu.get('degree', '')} | {edu.get('institution', '')}"
                    if edu.get("graduation_date"):
                        line += f" | {edu['graduation_date']}"
                    p = doc.add_paragraph(line)
                    p.runs[0].font.name = tmpl.font_name.value
                    p.runs[0].font.size = Pt(tmpl.font_size_body)
                continue

            content_map = {
                "summary": ("SUMMARY", data["summary"]),
                "skills": ("SKILLS", ", ".join(data["skills"]) if data["skills"] else ""),
                "certifications": ("CERTIFICATIONS", ", ".join(data["certifications"]) if data["certifications"] else ""),
            }
            header, content = content_map.get(section_key, ("", ""))
            if not content:
                continue
            if tmpl.include_section_headers:
                self._add_section_header(doc, header, tmpl)
            p = doc.add_paragraph(content)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = tmpl.line_spacing
            p.runs[0].font.name = tmpl.font_name.value
            p.runs[0].font.size = Pt(tmpl.font_size_body)

        doc.save(output_path)
        return output_path

    def _add_section_header(self, doc: Document, text: str, tmpl: ResumeTemplateSettings) -> None:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(tmpl.font_size_section)
        run.font.name = tmpl.font_name.value
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)

    def render_cover_letter_docx(self, body: str, output_path: Path) -> Path:
        tmpl = self.template
        doc = Document()
        for section in doc.sections:
            section.top_margin = Inches(tmpl.margin_inches)
            section.left_margin = Inches(tmpl.margin_inches)
            section.right_margin = Inches(tmpl.margin_inches)

        for para_text in body.split("\n\n"):
            if not para_text.strip():
                continue
            p = doc.add_paragraph(para_text.strip())
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            p.paragraph_format.line_spacing = 1.15
            if p.runs:
                p.runs[0].font.name = tmpl.font_name.value
                p.runs[0].font.size = Pt(tmpl.font_size_body)

        doc.save(output_path)
        return output_path

    def render_cover_letter_txt(self, body: str, output_path: Path) -> Path:
        output_path.write_text(body, encoding="utf-8")
        return output_path

    def render_linkedin_pack_docx(self, content: dict, output_path: Path) -> Path:
        tmpl = self.template
        doc = Document()
        sections = [
            ("HEADLINE OPTIONS", "\n\n".join(content.get("headline_variants", []))),
            ("OPTIMIZED ABOUT", content.get("optimized_about", "")),
            ("EXPERIENCE UPGRADES", self._format_upgrades(content.get("experience_upgrades", []))),
            ("SKILLS TO ADD", ", ".join(content.get("skills_to_add", []))),
        ]
        for header, body in sections:
            if not body:
                continue
            self._add_section_header(doc, header, tmpl)
            p = doc.add_paragraph(body)
            if p.runs:
                p.runs[0].font.name = tmpl.font_name.value
                p.runs[0].font.size = Pt(tmpl.font_size_body)
        doc.save(output_path)
        return output_path

    @staticmethod
    def _format_upgrades(upgrades: list) -> str:
        parts = []
        for item in upgrades:
            role = item.get("role", "Role")
            bullets = item.get("upgraded_bullets") or item.get("bullets") or str(item)
            parts.append(f"{role}\n{bullets}")
        return "\n\n".join(parts)

    def application_dir(self, application_id: str) -> Path:
        path = settings.applications_path / application_id
        path.mkdir(parents=True, exist_ok=True)
        return path


resume_renderer = ResumeRenderer()
