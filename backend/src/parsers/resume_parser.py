from pathlib import Path
import re
from typing import Optional

import pdfplumber
from docx import Document

from src.models.profile import CandidateProfile, ContactInfo, EducationEntry, ExperienceEntry
from src.scoring.text_normalize import normalize_resume_text


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(\+?\d[\d\s().-]{7,}\d)")
LINKEDIN_RE = re.compile(r"(https?://(?:www\.)?linkedin\.com/\S+)", re.I)
GITHUB_RE = re.compile(r"(https?://(?:www\.)?github\.com/\S+)", re.I)


def _sanitize_text(text: str) -> str:
    return normalize_resume_text(text)


def _parse_title_company_line(title_line: str) -> tuple[str, str, str | None]:
    company = "Unknown"
    title = title_line
    location = None
    if "|" in title_line:
        segments = [s.strip() for s in title_line.split("|")]
        title = segments[0]
        if len(segments) == 2:
            right = segments[1]
            if "," in right:
                location, company = [part.strip() for part in right.split(",", 1)]
            else:
                company = right
        elif len(segments) >= 3:
            location = segments[1]
            company = segments[2]
    elif " at " in title_line.lower():
        parts = re.split(r"\s+at\s+", title_line, maxsplit=1, flags=re.I)
        title, company = parts[0].strip(), parts[1].strip()
    return title, company, location


def _is_valid_experience_entry(title: str, company: str, bullets: list[str]) -> bool:
    if not title or title.lower().rstrip(".") in {"experience", "work experience", "employment"}:
        return False
    if company == "Unknown" and not bullets:
        return False
    return True


def extract_text_from_file(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf(file_path)
    elif suffix == ".docx":
        text = _extract_docx(file_path)
    elif suffix == ".txt":
        text = file_path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    return _sanitize_text(text)


def _extract_pdf(file_path: Path) -> str:
    text_parts: list[str] = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text)
    return "\n\n".join(text_parts)


def _extract_docx(file_path: Path) -> str:
    doc = Document(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def detect_layout_features(file_path: Path | None) -> dict[str, bool]:
    """Inspect DOCX structure for tables, multi-column sections, and embedded images."""
    result = {"has_tables": False, "has_multi_column": False, "has_images": False}
    if not file_path or not file_path.exists() or file_path.suffix.lower() != ".docx":
        return result
    try:
        doc = Document(file_path)
        body = doc.element.body
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        result["has_tables"] = bool(body.findall(".//w:tbl", ns))
        for sect in body.findall(".//w:sectPr", ns):
            cols = sect.find("w:cols", ns)
            if cols is not None:
                num = cols.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num")
                if num and int(num) > 1:
                    result["has_multi_column"] = True
        result["has_images"] = bool(body.findall(".//w:drawing", ns) or body.findall(".//w:pict", ns))
    except Exception:
        pass
    return result


def parse_profile_from_text(text: str, profile_id: str) -> CandidateProfile:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    contact = _extract_contact(text, lines)
    sections = _split_sections(text)
    skills = _extract_skills(sections.get("skills", ""))
    experience = _extract_experience(sections.get("experience", ""))
    education = _extract_education(sections.get("education", ""))
    summary = sections.get("summary") or sections.get("profile") or _first_paragraph(lines)

    return CandidateProfile(
        id=profile_id,
        contact=contact,
        summary=summary,
        skills=skills,
        experience=experience,
        education=education,
        resume_raw_text=text,
    )


def _extract_contact(text: str, lines: list[str]) -> ContactInfo:
    email_match = EMAIL_RE.search(text)
    phone_match = PHONE_RE.search(text)
    linkedin_match = LINKEDIN_RE.search(text)
    github_match = GITHUB_RE.search(text)
    name = lines[0] if lines else None
    return ContactInfo(
        name=name,
        email=email_match.group(0) if email_match else None,
        phone=phone_match.group(0).strip() if phone_match else None,
        linkedin_url=linkedin_match.group(1) if linkedin_match else None,
        github_url=github_match.group(1) if github_match else None,
    )


def _split_sections(text: str) -> dict[str, str]:
    section_headers = {
        "summary": r"^(summary|professional summary|profile|objective|about me|"
        r"résumé|resumen|profil|profilo|zusammenfassung|perfil|概要|简介)$",
        "experience": r"^(experience|work experience|employment|professional experience|"
        r"work history|career history|expérience|experiencia|experiência|erfahrung|"
        r"経歴|工作经历|employment history)$",
        "education": r"^(education|academic background|qualifications|"
        r"formation|educación|escolaridade|ausbildung|学歴|教育背景)$",
        "skills": r"^(skills|technical skills|core competencies|competencies|"
        r"compétences|habilidades|qualificações|fähigkeiten|スキル|技能)$",
        "certifications": r"^(certifications|licenses|certificates|"
        r"certificaciones|certificações|zertifikate)$",
    }
    lines = text.splitlines()
    sections: dict[str, list[str]] = {}
    current: Optional[str] = None

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        matched = False
        for key, pattern in section_headers.items():
            if re.match(pattern, line, re.I):
                current = key
                sections.setdefault(current, [])
                matched = True
                break
        if not matched and current:
            sections[current].append(line)

    return {key: "\n".join(value).strip() for key, value in sections.items()}


def _extract_skills(section: str) -> list[str]:
    if not section:
        return []
    parts = re.split(r"[,|•\n;]", section)
    return [part.strip() for part in parts if part.strip()][:40]


def _extract_experience(section: str) -> list[ExperienceEntry]:
    if not section:
        return []
    lines = [line.strip() for line in section.splitlines() if line.strip()]
    entries: list[ExperienceEntry] = []
    i = 0
    date_pattern = re.compile(
        r"^(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|"
        r"January|February|March|April|June|July|August|September|October|November|December|"
        r"Janvier|Févr|Février|Mars|Avr|Avril|Mai|Juin|Juil|Juillet|Août|Sept|Oct|Nov|Déc|"
        r"Ene|Feb|Mar|Abr|May|Jun|Jul|Ago|Sep|Oct|Nov|Dic)[a-zéûî\.]*\.?\s+\d{4}"
        r"|\d{1,2}[/.-]\d{4}|\d{4}[/.-]\d{1,2}|\d{4})\s*"
        r"[—\-–~to]+\s*"
        r"(?:(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|"
        r"January|February|March|April|June|July|August|September|October|November|December|"
        r"Janvier|Févr|Février|Mars|Avr|Avril|Mai|Juin|Juil|Juillet|Août|Sept|Oct|Nov|Déc|"
        r"Ene|Feb|Mar|Abr|May|Jun|Jul|Ago|Sep|Oct|Nov|Dic)[a-zéûî\.]*\.?\s+\d{4}"
        r"|\d{1,2}[/.-]\d{4}|\d{4}[/.-]\d{1,2}|\d{4}|Current|Present|Now|Presente|Actuel|Heute|今)",
        re.I,
    )

    while i < len(lines):
        line = lines[i]
        start_date, end_date = None, None
        title_line = line

        if date_pattern.match(line):
            parts = re.split(r"\s*[—\-–]\s*", line, maxsplit=1)
            start_date = parts[0].strip() if parts else None
            end_date = parts[1].strip() if len(parts) > 1 else None
            i += 1
            if i >= len(lines):
                break
            title_line = lines[i]

        company = "Unknown"
        title = title_line
        location = None
        title, company, location = _parse_title_company_line(title_line)

        bullets: list[str] = []
        i += 1
        while i < len(lines):
            next_line = lines[i]
            if date_pattern.match(next_line):
                break
            if re.match(r"^(EXPERIENCE|EDUCATION|SKILLS|CERTIFICATIONS|SUMMARY|WEBSITES)", next_line, re.I):
                break
            if re.match(r"^(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)", next_line, re.I) and "—" in next_line:
                break
            if next_line.startswith(("-", "•", "*")):
                bullets.append(next_line.lstrip("-•* ").strip())
            elif len(next_line) > 20 and not re.match(r"^[A-Z][a-z]+ [A-Z]", next_line):
                bullets.append(next_line)
            elif "|" in next_line and i > 0:
                break
            else:
                if len(next_line) > 30:
                    bullets.append(next_line)
            i += 1

        if not _is_valid_experience_entry(title, company, bullets):
            continue

        entries.append(
            ExperienceEntry(
                title=title,
                company=company,
                location=location,
                start_date=start_date,
                end_date=end_date,
                bullets=bullets[:8],
            )
        )

    if entries:
        return entries[:10]

    blocks = re.split(r"\n(?=[A-Z][^\n]{0,80}(?:\|| at | - ))", section)
    for block in blocks:
        blines = [line.strip() for line in block.splitlines() if line.strip()]
        if not blines:
            continue
        title_line = blines[0]
        company = "Unknown"
        title = title_line
        if "|" in title_line:
            title, company = [part.strip() for part in title_line.split("|", 1)]
        elif " at " in title_line.lower():
            parts = re.split(r"\s+at\s+", title_line, maxsplit=1, flags=re.I)
            title, company = parts[0].strip(), parts[1].strip()
        bullets = [line.lstrip("-•* ").strip() for line in blines[1:] if line.startswith(("-", "•", "*"))]
        entries.append(ExperienceEntry(title=title, company=company, bullets=bullets))
    return entries[:10]


def _extract_education(section: str) -> list[EducationEntry]:
    if not section:
        return []
    entries: list[EducationEntry] = []
    for line in section.splitlines():
        line = line.strip()
        if not line:
            continue
        if "|" in line:
            degree, institution = [part.strip() for part in line.split("|", 1)]
        else:
            degree, institution = line, "Institution"
        entries.append(EducationEntry(degree=degree, institution=institution))
    return entries[:5]


def _first_paragraph(lines: list[str]) -> Optional[str]:
    if len(lines) > 1:
        return lines[1]
    return None
